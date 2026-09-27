"""
Builds the Research Archive of Rising Scholars (RARS) version of the paper:
manuscript/rars/paper_rars.docx (from the RARS Word template) and, through
Microsoft Word, manuscript/rars/paper_rars.pdf.

paper.tex is the single source. This script converts the subset of LaTeX the
paper uses (sections, paragraphs, emphasis, citations, cross-references,
inline math, tables fed by tables/*.tex, figures) and resolves every number
macro from numbers.tex, so the Word version carries exactly the same
generated values. Citations and references are rendered in APA 7 style;
the reference list is numbered for Google Scholar indexing, as RARS asks.
The references come from analysis/apa_refs.py, which also writes the LaTeX
reference lists, so both versions carry the same entries. References to the
Supplementary Information (\ref{S-...}) are numbered from supplement.aux, so
build the LaTeX documents first. Body text is Arial 12 pt, single-spaced;
tables are set in Arial 10 pt; figures are embedded at 300 dpi.

The RARS Word template is not redistributed; point RARS_TEMPLATE at a copy
(default: "RARS template.docx" next to the manuscript directory).

Run from the manuscript directory:  python analysis/make_docx.py
"""
import os
import re
import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
MS = HERE.parent
OUT = MS / "rars"
OUT.mkdir(exist_ok=True)
TEMPLATE = Path(os.environ.get("RARS_TEMPLATE", MS.parent / "RARS template.docx"))
FONT = "Arial"
TEXT_W = 8.5 - 2 * 0.7  # inches, template margins
MINUS = "\u2212"


# ------------------------------------------------------------------ brace utilities
def group(s, i):
    """s[i] == '{': return (content, index after the closing brace)."""
    assert s[i] == "{", s[i:i + 30]
    depth, j = 0, i
    while True:
        c = s[j]
        if c == "\\":
            j += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1


def opt(s, i):
    """s[i] == '[': return (content, index after ']')."""
    depth, j = 0, i
    while True:
        if s[j] == "[":
            depth += 1
        elif s[j] == "]":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1


# ------------------------------------------------------------------ macros
def load_macros():
    src = (MS / "numbers.tex").read_text(encoding="utf-8")
    macros = {}
    for m in re.finditer(r"\\newcommand\{\\([A-Za-z]+)\}", src):
        val, _ = group(src, m.end())
        macros[m.group(1)] = val
    return macros


MACROS = load_macros()
MACROS.update({"HK": "H\\&K", "pp": "\\,pp", "SDu": "\\,SD", "sh": "\u266f"})


def expand(s):
    """Expand number macros (they contain no arguments)."""
    for _ in range(4):
        s2 = re.sub(r"\\([A-Za-z]+)(\{\})?",
                    lambda m: MACROS[m.group(1)] if m.group(1) in MACROS else m.group(0), s)
        if s2 == s:
            break
        s = s2
    return s


# ------------------------------------------------------------------ bibliography (APA 7)
sys.path.insert(0, str(HERE))
from apa_refs import BIB, cite_names, ref_entry, ref_sort_key, year_label  # noqa: E402

CITED = []


def note_cite(k):
    if k not in BIB:
        raise KeyError(f"unknown citation key {k}")
    if k not in CITED:
        CITED.append(k)


def citep(keys, pre="", post=""):
    items = []  # [names, [years]]; consecutive works by the same authors share the names
    for k in keys:
        note_cite(k)
        names = cite_names(k, False)
        if items and items[-1][0] == names:
            items[-1][1].append(year_label(k))
        else:
            items.append([names, [year_label(k)]])
    body = "; ".join(f"{n}, {', '.join(y)}" for n, y in items)
    if pre:
        body = pre + " " + body
    if post:
        body = body + ", " + post
    return "(" + body + ")"


def citet(keys):
    items = []
    for k in keys:
        note_cite(k)
        names = cite_names(k, True)
        if items and items[-1][0] == names:
            items[-1][1].append(year_label(k))
        else:
            items.append([names, [year_label(k)]])
    return ", ".join(f"{n} ({', '.join(y)})" for n, y in items)


def citealp(keys):
    out = []
    for k in keys:
        note_cite(k)
        out.append(f"{cite_names(k, False)}, {year_label(k)}")
    return "; ".join(out)


# ------------------------------------------------------------------ inline conversion
GREEK = {"alpha": "\u03b1", "beta": "\u03b2", "sigma": "\u03c3", "tau": "\u03c4", "lambda": "\u03bb",
         "Delta": "\u0394", "rho": "\u03c1", "mu": "\u03bc", "chi": "\u03c7", "pi": "\u03c0"}
SYMS = {"ge": "\u2265", "le": "\u2264", "geq": "\u2265", "leq": "\u2264", "times": "\u00d7", "pm": "\u00b1",
        "approx": "\u2248", "cdot": "\u00b7", "infty": "\u221e", "sharp": "\u266f", "dagger": "\u2020",
        "ddagger": "\u2021", "to": "\u2192", "neq": "\u2260", "log": "log", "min": "min", "max": "max",
        "exp": "exp", "ln": "ln", "cos": "cos", "sum": "\u2211", "quad": "\u2003", "sim": "~", "prime": "\u2032"}


class Run:
    __slots__ = ("text", "b", "i", "sup", "sub")

    def __init__(self, text, b=False, i=False, sup=False, sub=False):
        self.text, self.b, self.i, self.sup, self.sub = text, b, i, sup, sub


def math_runs(s, fmt):
    """Minimal TeX math -> styled runs. Letters italic, digits upright."""
    s = expand(s)
    out = []
    i = 0

    def emit(t, **kw):
        f = dict(fmt)
        f.update(kw)
        out.append(Run(t, **f))

    while i < len(s):
        c = s[i]
        if c == "\\":
            m = re.match(r"\\([A-Za-z]+|.)", s[i:])
            name = m.group(1)
            i += len(m.group(0))
            if name in GREEK:
                emit(GREEK[name], i=True)
            elif name in SYMS:
                emit(SYMS[name])
            elif name in ("mathrm", "text", "textrm", "operatorname", "mathit"):
                while s[i] == " ":
                    i += 1
                g, i = group(s, i)
                sub_runs = inline_runs(g, fmt) if name == "text" else [Run(expand(g), **fmt)]
                out.extend(sub_runs)
            elif name in ("tfrac", "frac", "dfrac"):
                if s[i] == "{":
                    a, i = group(s, i)
                else:
                    a, i = s[i], i + 1
                if s[i] == "{":
                    b, i = group(s, i)
                else:
                    b, i = s[i], i + 1
                if (a, b) == ("1", "2"):
                    emit("\u00bd")
                else:
                    out.extend(math_runs(a, fmt))
                    emit("/")
                    out.extend(math_runs(b, fmt))
            elif name == "sqrt":
                g, i = group(s, i)
                emit("\u221a(")
                out.extend(math_runs(g, fmt))
                emit(")")
            elif name == "bar":
                while s[i] == " ":
                    i += 1
                if s[i] == "{":
                    g, i = group(s, i)
                else:
                    g, i = s[i], i + 1
                emit(g + "\u0304", i=True)
            elif name in ("bigl", "bigr", "left", "right", "Bigl", "Bigr"):
                pass
            elif name == ",":
                emit("\u2009")
            elif name in ("%", "&", "{", "}", "_", "#"):
                emit(name)
            elif name == "textminus":
                emit(MINUS)
            else:
                raise ValueError(f"unhandled math command \\{name} in: {s}")
        elif c in "^_":
            i += 1
            if s[i] == "{":
                g, i = group(s, i)
            elif s[i] == "\\":
                m = re.match(r"\\[A-Za-z]+", s[i:])
                g, i = m.group(0), i + len(m.group(0))
            else:
                g, i = s[i], i + 1
            sub_fmt = dict(fmt)
            sub_fmt["sup" if c == "^" else "sub"] = True
            out.extend(math_runs(g, sub_fmt))
        elif c == "-":
            emit(MINUS)
            i += 1
        elif c in "{}" or c.isspace():  # spaces in math are not typeset
            i += 1
        elif c == "~":
            emit("\u00a0")
            i += 1
        elif c.isalpha():
            emit(c, i=True)
            i += 1
        elif c in "=<>+":
            emit(f"\u2009{c}\u2009" if c in "=<>" else c)
            i += 1
        else:
            emit(c)
            i += 1
    return out


def text_subst(t):
    t = t.replace("---", "\u2014").replace("--", "\u2013")
    t = t.replace("``", "\u201c").replace("''", "\u201d").replace("`", "\u2018").replace("'", "\u2019")
    t = t.replace("~", "\u00a0")
    return t


def inline_runs(s, fmt=None):
    """Convert LaTeX inline text to styled runs."""
    fmt = fmt or {"b": False, "i": False, "sup": False, "sub": False}
    s = expand(s)
    out = []
    buf = []
    i = 0

    def flush():
        if buf:
            out.append(Run(text_subst("".join(buf)), **fmt))
            buf.clear()

    while i < len(s):
        c = s[i]
        if c == "%":  # comment to end of line
            j = s.find("\n", i)
            i = len(s) if j < 0 else j + 1
            continue
        if c == "$":
            flush()
            j = s.index("$", i + 1)
            out.extend(math_runs(s[i + 1:j], fmt))
            i = j + 1
            continue
        if c in "{}":
            i += 1
            continue
        if c == "\n":
            buf.append(" ")
            i += 1
            continue
        if c != "\\":
            buf.append(c)
            i += 1
            continue
        m = re.match(r"\\([A-Za-z]+|.)", s[i:], re.DOTALL)
        name = m.group(1)
        i += len(m.group(0))
        if name in (" ", "\n"):
            buf.append(" ")
            continue
        if name in ("%", "&", "#", "_", "{", "}"):
            buf.append(name)
            continue
        if name == ",":
            buf.append("\u2009")
            continue
        if name == " ":
            buf.append(" ")
            continue
        if name == "\\":
            buf.append(" ")
            continue
        if name in ("textminus",):
            buf.append(MINUS)
            continue
        if name == "texttimes":
            buf.append("\u00d7")
            continue
        if name == "ldots":
            buf.append("\u2026")
            continue
        if name == "quad":
            buf.append("\u2003")
            continue
        if name == "sharp":
            buf.append("\u266f")
            continue
        if name in ("emph", "textit", "textbf", "texttt", "textsc", "mbox", "flip", "textsuperscript"):
            g, i = group(s, i)
            flush()
            f = dict(fmt)
            if name in ("emph", "textit"):
                f["i"] = not fmt["i"]
            if name in ("textbf", "flip"):
                f["b"] = True
            if name == "textsuperscript":
                f["sup"] = True
            out.extend(inline_runs(g, f))
            continue
        if name in ("citep", "citet", "citealp", "cite"):
            opts = []
            while i < len(s) and s[i] == "[":
                o, i = opt(s, i)
                opts.append(o)
            g, i = group(s, i)
            keys = [k.strip() for k in g.split(",")]
            if name == "citep" or name == "cite":
                pre, post = (opts + ["", ""])[:2] if len(opts) == 2 else ("", opts[0] if opts else "")
                buf.append(citep(keys, text_subst(pre), text_subst(post)))
            elif name == "citet":
                buf.append(citet(keys))
            else:
                buf.append(citealp(keys))
            continue
        if name == "ref":
            g, i = group(s, i)
            buf.append(SUPP_LABELS[g[2:]] if g.startswith("S-") else LABELS[g])
            continue
        if name in ("label", "index"):
            _, i = group(s, i)
            continue
        if name in ("noindent", "medskip", "smallskip", "bigskip", "centering", "small", "footnotesize",
                    "normalsize", "par", "hfill", "relax", "protect"):
            continue
        if name == "url":
            g, i = group(s, i)
            buf.append(g)
            continue
        raise ValueError(f"unhandled command \\{name} near: {s[max(0, i - 60):i + 20]!r}")
    flush()
    return out


# ------------------------------------------------------------------ document parsing
TEX = (MS / "paper.tex").read_text(encoding="utf-8")
BODY = TEX[TEX.index("\\begin{document}") + len("\\begin{document}"):TEX.index("\\end{document}")]
BODY = re.sub(r"(?m)^%.*\n", "", BODY)
BODY = re.sub(r"(?<!\\)%[^\n]*", "", BODY)

LABELS = {}


def load_supp_labels():
    aux = MS / "supplement.aux"
    if not aux.exists():
        sys.exit("supplement.aux not found: build supplement.tex first")
    out = {}
    for m in re.finditer(r"\\newlabel\{([^}]*)\}\{\{([^}]*)\}", aux.read_text(encoding="utf-8")):
        out[m.group(1)] = m.group(2)
    return out


SUPP_LABELS = load_supp_labels()
AUTHOR_TEX, _ = group(TEX, TEX.index("\\author{") + len("\\author"))
AUTHOR, AFFIL = [re.sub(r"\\normalsize\s*", "", x).strip()
                 for x in re.split(r"\\\\(?:\[[^\]]*\])?", AUTHOR_TEX)]


def number_labels():
    sec = sub = tab = fig = 0
    appendix = False
    current = None
    for m in re.finditer(r"\\(section|subsection|appendix|label|begin)\*?(\{[^}]*\})?", BODY):
        kind = m.group(1)
        arg = (m.group(2) or "")[1:-1]
        star = BODY[m.start():m.end()].startswith(("\\section*", "\\subsection*"))
        if kind == "appendix":
            appendix, sec = True, 0
        elif kind == "section" and not star:
            sec += 1
            sub = 0
            current = "ABCDEFG"[sec - 1] if appendix else str(sec)
        elif kind == "subsection" and not star:
            sub += 1
            current = f"{'ABCDEFG'[sec - 1] if appendix else sec}.{sub}"
        elif kind == "begin" and arg == "table":
            tab += 1
            current = str(tab)
        elif kind == "begin" and arg == "figure":
            fig += 1
            current = str(fig)
        elif kind == "label":
            LABELS[arg] = current


number_labels()


# ------------------------------------------------------------------ docx helpers
def set_run_font(run, size=12):
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT)


def fmt_par(p, align=None, after=6, before=0, keep=False, indent=None, hanging=None):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    if align is not None:
        p.alignment = align
    if keep:
        pf.keep_with_next = True
    if indent is not None:
        pf.left_indent = Inches(indent)
    if hanging is not None:
        pf.first_line_indent = Inches(-hanging)


def add_runs(p, runs, size=12, bold=False, italic=False):
    first = not p.runs
    for r in runs:
        text = r.text.lstrip() if first else r.text
        if not text:
            continue
        first = False
        run = p.add_run(text)
        set_run_font(run, size)
        run.bold = r.b or bold
        run.italic = r.i != italic if italic else r.i
        run.font.superscript = r.sup
        run.font.subscript = r.sub


def para(doc, tex, size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY, lead=None, after=6):
    p = doc.add_paragraph()
    if "\\url" in tex and align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        align = WD_ALIGN_PARAGRAPH.LEFT  # long URLs would stretch justified lines
    fmt_par(p, align, after=after)
    if lead:
        add_runs(p, inline_runs(lead), size, bold=True)
        add_runs(p, [Run(" ")], size)
    add_runs(p, inline_runs(tex.strip()), size)
    return p


def heading(doc, text, level):
    p = doc.add_paragraph()
    fmt_par(p, WD_ALIGN_PARAGRAPH.LEFT, after=6, before=12 if level == 1 else 8, keep=True)
    add_runs(p, inline_runs(text), 12, bold=True, italic=(level == 2))
    return p


def set_cell_border(cell, **edges):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, val in edges.items():
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single" if val else "nil")
        if val:
            el.set(qn("w:sz"), str(val))
            el.set(qn("w:color"), "000000")
        borders.append(el)


def split_cells(line):
    cells, cur, depth, i = [], [], 0, 0
    while i < len(line):
        c = line[i]
        if c == "\\" and i + 1 < len(line):
            cur.append(line[i:i + 2])
            i += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
        if c == "&" and depth == 0:
            cells.append("".join(cur))
            cur = []
        else:
            cur.append(c)
        i += 1
    cells.append("".join(cur))
    return [c.strip() for c in cells]


def parse_row(line):
    """-> list of (text, colspan)."""
    out = []
    for c in split_cells(line):
        m = re.match(r"\\multicolumn\{(\d+)\}\{[^}]*\}", c)
        if m:
            g, _ = group(c, m.end())
            out.append((g, int(m.group(1))))
        else:
            out.append((c, 1))
    return out


def table_rows(tab):
    """Tabular body -> (header rows, body rows); rows are lists of (tex, span)."""
    tab = re.sub(r"\\tabinput\{([^}]*)\}", lambda m: (MS / m.group(1)).read_text(encoding="utf-8"), tab)
    tab = re.sub(r"\\cmidrule(\([^)]*\))?\{[^}]*\}", "", tab)
    head, _, body = tab.partition("\\midrule")
    head = head.replace("\\toprule", "")
    body = body.replace("\\bottomrule", "")
    body = re.sub(r"\\addlinespace(\[[^\]]*\])?|\\midrule", "", body)
    split = lambda t: [r.strip() for r in re.split(r"\\\\(?:\[[^\]]*\])?", t) if r.strip()]
    return [parse_row(r) for r in split(head)], [parse_row(r) for r in split(body)]


def cell_text(tex):
    return "".join(r.text for r in inline_runs(tex))


def column_widths(spec, rows, ncol, size):
    """Column widths (inches) from the longest body entry of each column; header text may
    wrap at word boundaries. Returns (widths, font size)."""
    fixed = {}
    for n, m in enumerate(re.finditer(r"p\{([0-9.]*)\\linewidth\}|[lcr]", spec)):
        if m.group(1) is not None:
            fixed[n] = float(m.group(1) or 1) * TEXT_W
    need = [0.0] * ncol
    for row, is_head in rows:
        col = 0
        for tex, span in row:
            if col >= ncol:
                break
            if span == 1:
                t = cell_text(tex).strip()
                n = max(len(w) for w in t.split()) if (is_head and t) else len(t)
                need[col] = max(need[col], n)
            col += span
    for size in (size, 9, 8):
        char = 0.52 * size / 72  # average Arial glyph width, inches
        w = [fixed.get(c, need[c] * char + 0.14) for c in range(ncol)]
        if sum(w) <= TEXT_W:
            break
    if sum(w) > TEXT_W:
        # wrap the first (label) column first, then shrink everything proportionally
        over = sum(w) - TEXT_W
        w[0] = max(w[0] - over, 0.9)
        if sum(w) > TEXT_W:
            w = [x * TEXT_W / sum(w) for x in w]
    else:
        # spread the spare width over the data columns
        spare = TEXT_W - sum(w)
        data = list(range(1, ncol)) or [0]
        for c in data:
            w[c] += spare / len(data) * 0.5
    return w, size


def set_table_props(t, widths):
    tbl = t._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", 0), ("bottom", 0), ("left", 60), ("right", 60)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)
    grid = tbl.tblGrid
    for gc, wd in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(wd * 1440)))


def mark_row(row, header):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)
    if header:
        el = OxmlElement("w:tblHeader")
        trPr.append(el)


def add_table(doc, caption, label_num, spec, tab, size=10):
    heads, body = table_rows(tab)
    ncol = len(re.findall(r"[lcrp]", re.sub(r"p\{[^}]*\}", "p", spec)))
    all_rows = [(r, True) for r in heads] + [(r, False) for r in body]
    widths, size = column_widths(spec, all_rows, ncol, size)
    cap = doc.add_paragraph()
    fmt_par(cap, WD_ALIGN_PARAGRAPH.LEFT, after=4, before=6, keep=True)
    add_runs(cap, [Run(f"Table {label_num}. ", b=True)], 10)
    add_runs(cap, inline_runs(caption), 10)
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_props(t, widths)
    aligns = [a for a in re.findall(r"[lcrp]", re.sub(r"p\{[^}]*\}", "p", spec))]
    for n, (row, is_head) in enumerate(all_rows):
        r = t.add_row()
        mark_row(r, is_head)
        cells = r.cells
        for c, wd in zip(cells, widths):
            c.width = Inches(wd)
        col = 0
        for tex, span in row:
            if col >= ncol:
                break
            cell = cells[col]
            if span > 1:
                cell = cell.merge(cells[min(col + span - 1, ncol - 1)])
            cell.text = ""
            p = cell.paragraphs[0]
            fmt_par(p, {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                        "r": WD_ALIGN_PARAGRAPH.RIGHT, "p": WD_ALIGN_PARAGRAPH.LEFT}[aligns[col]]
                    if span == 1 else WD_ALIGN_PARAGRAPH.CENTER, after=0, before=1)
            tex = tex.strip()
            if tex.startswith("\\quad"):
                tex = tex[len("\\quad"):]
                p.paragraph_format.left_indent = Inches(0.15)
            if n < len(all_rows) - 1:
                p.paragraph_format.keep_with_next = is_head
            add_runs(p, inline_runs(tex), size, bold=False)
            col += span
        # rules: top of first row, bottom of last header row, bottom of last row
        for c in t.rows[-1].cells:
            edges = {}
            if n == 0:
                edges["top"] = 8
            if is_head and n == len(heads) - 1:
                edges["bottom"] = 4
            if n == len(all_rows) - 1:
                edges["bottom"] = 8
            if edges:
                set_cell_border(c, **edges)
    after = doc.add_paragraph()
    fmt_par(after, after=6)


def add_figure(doc, path_tex, width_frac, caption, num):
    import pymupdf
    src = MS / path_tex
    png = OUT / (src.stem + ".png")
    d = pymupdf.open(src)
    d[0].get_pixmap(dpi=300).save(png)
    p = doc.add_paragraph()
    fmt_par(p, WD_ALIGN_PARAGRAPH.CENTER, after=2, before=6, keep=True)
    p.add_run().add_picture(str(png), width=Inches(TEXT_W * width_frac))
    cap = doc.add_paragraph()
    fmt_par(cap, WD_ALIGN_PARAGRAPH.JUSTIFY, after=8)
    add_runs(cap, [Run(f"Figure {num}. ", b=True)], 10)
    add_runs(cap, inline_runs(caption), 10)


# ------------------------------------------------------------------ main conversion
def build():
    doc = docx.Document(str(TEMPLATE))
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]

    title_tex = re.search(r"\\title\{", TEX)
    title, _ = group(TEX, title_tex.end() - 1)
    p = doc.add_paragraph()
    fmt_par(p, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_runs(p, inline_runs(title.replace("\\\\", " ")), 12, bold=True)
    for line, after in ((AUTHOR, 0), (AFFIL, 12)):
        p = doc.add_paragraph()
        fmt_par(p, WD_ALIGN_PARAGRAPH.CENTER, after=after)
        add_runs(p, [Run(line)], 12)

    s = BODY.replace("\\maketitle", "")
    # abstract
    a0, a1 = s.index("\\begin{abstract}"), s.index("\\end{abstract}")
    abstract = s[a0 + len("\\begin{abstract}"):a1]
    s = s[:a0] + s[a1 + len("\\end{abstract}"):]
    heading(doc, "Abstract", 1)
    abs_text, _, kw = abstract.partition("\\medskip")
    para(doc, abs_text)
    kw = re.sub(r"\\noindent\\textbf\{Keywords:\}", "", kw)
    para(doc, kw, lead="Keywords:")
    global ABSTRACT_WORDS, TITLE_TEXT
    runs = inline_runs(abs_text)
    plain = "".join(r.text if not r.sup else ("²" if r.text == "2" else "^" + r.text) for r in runs)
    plain = re.sub(r"\s+", " ", plain).strip()
    ABSTRACT_WORDS = len(plain.split())
    (OUT / "abstract.txt").write_text(plain + "\n", encoding="utf-8")
    TITLE_TEXT = " ".join("".join(r.text for r in inline_runs(title.replace("\\\\", " "))).split())

    tokens = re.split(r"(\\section\*?\{|\\subsection\*?\{|\\paragraph\{|\\begin\{table\}(?:\[[^\]]*\])?|"
                      r"\\begin\{figure\}(?:\[[^\]]*\])?|\\begin\{enumerate\}|\\appendix|"
                      r"\\input\{bib_paper(?:\.tex)?\})", s)
    sec = sub = 0
    appendix = False
    tab_n = fig_n = 0
    pending_lead = None
    i = 0

    def paragraphs(chunk, lead=None):
        # skip blocks that typeset nothing (a lone \label after a heading, for example)
        blocks = [b for b in re.split(r"\n\s*\n", chunk)
                  if b.strip() and "".join(r.text for r in inline_runs(b)).strip()]
        for n, b in enumerate(blocks):
            para(doc, b, lead=lead if n == 0 else None)
        return bool(blocks)

    while i < len(tokens):
        tk = tokens[i]
        i += 1
        if tk.startswith("\\section") or tk.startswith("\\subsection"):
            star = "*" in tk
            rest = "{" + tokens[i]
            name, j = group(rest, 0)
            tokens[i] = rest[j:]
            if tk.startswith("\\section"):
                if not star:
                    sec += 1
                    sub = 0
                num = ("ABCDEFG"[sec - 1] if appendix else str(sec)) if not star else ""
                heading(doc, (f"Appendix {num}. " if appendix and num else (f"{num} " if num else "")) + name, 1)
            else:
                sub += 1
                num = f"{'ABCDEFG'[sec - 1] if appendix else sec}.{sub}"
                heading(doc, f"{num} {name}", 2)
            continue
        if tk == "\\paragraph{":
            rest = "{" + tokens[i]
            name, j = group(rest, 0)
            lead = name if name.endswith((".", "?", "!")) else name + "."
            chunk = rest[j:]
            blocks = [b for b in re.split(r"\n\s*\n", chunk) if b.strip()]
            for n, b in enumerate(blocks):
                para(doc, b, lead=lead if n == 0 else None)
            tokens[i] = ""
            continue
        if tk.startswith("\\begin{table}"):
            env, _, rest = tokens[i].partition("\\end{table}")
            tokens[i] = rest
            tab_n += 1
            cap_i = env.index("\\caption{")
            cap, _ = group(env, cap_i + len("\\caption"))
            tb = env.index("\\begin{tabular}")
            spec, k = group(env, tb + len("\\begin{tabular}"))
            tab = env[k:env.index("\\end{tabular}")]
            add_table(doc, cap, tab_n, spec, tab)
            continue
        if tk.startswith("\\begin{figure}"):
            env, _, rest = tokens[i].partition("\\end{figure}")
            tokens[i] = rest
            fig_n += 1
            m = re.search(r"\\includegraphics\[width=([0-9.]*)\\linewidth\]\{([^}]*)\}", env)
            frac = float(m.group(1) or 1)
            cap, _ = group(env, env.index("\\caption{") + len("\\caption"))
            add_figure(doc, m.group(2), frac, cap, fig_n)
            continue
        if tk == "\\begin{enumerate}":
            env, _, rest = tokens[i].partition("\\end{enumerate}")
            tokens[i] = rest
            env = re.sub(r"\\itemsep\s*\d+pt", "", env)
            for item in re.split(r"\\item", env)[1:]:
                label = ""
                item = item.strip()
                if item.startswith("["):
                    label, k = opt(item, 0)
                    item = item[k:]
                p = para(doc, item, lead=None)
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.first_line_indent = Inches(-0.4)
                if label:
                    r0 = p.runs[0]
                    r0.text = label + "\t" + r0.text
            continue
        if tk == "\\appendix":
            appendix, sec = True, 0
            continue
        if tk.startswith("\\input{bib_paper"):
            heading(doc, "References", 1)
            for n, k in enumerate(sorted(CITED, key=ref_sort_key), 1):
                p = doc.add_paragraph()
                fmt_par(p, WD_ALIGN_PARAGRAPH.LEFT, after=4, indent=0.4, hanging=0.4)
                add_runs(p, [Run(f"{n}.\t")] + [Run(t, i=it) for t, it in ref_entry(k)], 12)
            continue
        paragraphs(tk)
    return doc


def export_pdf(docx_path, pdf_path):
    import win32com.client
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(str(docx_path), ReadOnly=True)
        d.ExportAsFixedFormat(str(pdf_path), 17)
        d.Close(False)
    finally:
        word.Quit()


if __name__ == "__main__":
    doc = build()
    cp = doc.core_properties
    cp.author, cp.last_modified_by, cp.title = AUTHOR, AUTHOR, TITLE_TEXT
    cp.subject = cp.keywords = cp.comments = cp.category = ""
    path = OUT / "paper_rars.docx"
    doc.save(path)
    print("wrote", path, "abstract words:", ABSTRACT_WORDS, "references:", len(CITED))
    if "--no-pdf" not in sys.argv:
        export_pdf(path, OUT / "paper_rars.pdf")
        import pymupdf
        pdf = pymupdf.open(OUT / "paper_rars.pdf")
        pdf.set_metadata({**pdf.metadata, "title": TITLE_TEXT, "author": AUTHOR})
        pdf.saveIncr()
        print("wrote", OUT / "paper_rars.pdf")
